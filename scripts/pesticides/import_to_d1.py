#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
FAMIC 農薬登録情報 + CropLife RAC 製品名順表 → ローカル SQLite / D1 投入

使い方:
  python scripts/pesticides/import_to_d1.py --input-dir "C:/path/to/20261019"
  python scripts/pesticides/import_to_d1.py --input-dir "..." --apply-d1-local

期待する配置（どれでも可）:
  <dir>/登録基本部.xls
  <dir>/登録適用部一.xls
  <dir>/登録適用部二.xls
  <dir>/mechanism_rac.xlsx

  またはサブフォルダ:
  <dir>/登録基本部/登録基本部.xls
  <dir>/登録適用部一/登録適用部一.xls
  <dir>/登録適用部二/登録適用部二.xls
"""

from __future__ import annotations

import argparse
import json
import os
import shutil
import sqlite3
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

try:
    import xlrd
except ImportError:
    print("xlrd が必要です: pip install -r scripts/pesticides/requirements.txt", file=sys.stderr)
    raise

try:
    from openpyxl import load_workbook
except ImportError:
    print("openpyxl が必要です: pip install -r scripts/pesticides/requirements.txt", file=sys.stderr)
    raise

ROOT = Path(__file__).resolve().parents[2]
MIGRATION = ROOT / "migrations" / "0001_pesticides.sql"
OUT_DIR = ROOT / "data" / "pesticides"
SQLITE_PATH = OUT_DIR / "local.sqlite"
SQL_BATCH_DIR = OUT_DIR / "d1_batches"
DEFAULT_DB_NAME = "tool-portal-pesticides"


def find_file(input_dir: Path, *candidates: str) -> Path:
    for name in candidates:
        direct = input_dir / name
        if direct.is_file():
            return direct
        nested = input_dir / name.replace(".xls", "") / name
        if nested.is_file():
            return nested
        # folder name may equal stem
        stem = Path(name).stem
        nested2 = input_dir / stem / name
        if nested2.is_file():
            return nested2
    raise FileNotFoundError(f"見つかりません: {candidates} (in {input_dir})")


def cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if v == int(v):
            return str(int(v))
        return str(v)
    return str(v).strip()


def load_xls_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    book = xlrd.open_workbook(str(path))
    sh = book.sheet_by_index(0)
    headers = [cell_str(sh.cell_value(0, c)) for c in range(sh.ncols)]
    rows: list[dict[str, str]] = []
    for r in range(1, sh.nrows):
        row = {headers[c]: cell_str(sh.cell_value(r, c)) for c in range(sh.ncols)}
        rows.append(row)
    return headers, rows


def sql_quote(value: str | None) -> str:
    if value is None:
        return "NULL"
    return "'" + str(value).replace("'", "''") + "'"


def init_sqlite(db_path: Path) -> sqlite3.Connection:
    if db_path.exists():
        db_path.unlink()
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(db_path))
    conn.execute("PRAGMA foreign_keys = ON")
    conn.executescript(MIGRATION.read_text(encoding="utf-8"))
    return conn


def import_famic_basic(conn: sqlite3.Connection, path: Path) -> tuple[int, int]:
    _, rows = load_xls_rows(path)
    products: dict[str, dict[str, str]] = {}
    ingredients: list[tuple[str, str, str]] = []

    for row in rows:
        reg = row.get("登録番号", "")
        if not reg:
            continue
        if reg not in products:
            mix = row.get("混合数", "")
            products[reg] = {
                "reg_no": reg,
                "pesticide_type": row.get("農薬の種類", ""),
                "name": row.get("農薬の名称", "") or "(名称なし)",
                "holder": row.get("登録を有する者の名称", ""),
                "concentration": row.get("濃度", ""),
                "mix_count": mix,
                "use_category": row.get("用途", ""),
                "formulation": row.get("剤型名", ""),
                "registered_on": row.get("登録年月日", ""),
            }
        ing = row.get("有効成分", "")
        if ing:
            ingredients.append((reg, ing, row.get("総使用回数における有効成分", "")))

    cur = conn.cursor()
    cur.executemany(
        """
        INSERT INTO products (
          reg_no, pesticide_type, name, holder, concentration,
          mix_count, use_category, formulation, registered_on
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [
            (
                p["reg_no"],
                p["pesticide_type"],
                p["name"],
                p["holder"],
                p["concentration"],
                int(p["mix_count"]) if p["mix_count"].isdigit() else None,
                p["use_category"],
                p["formulation"],
                p["registered_on"],
            )
            for p in products.values()
        ],
    )
    cur.executemany(
        """
        INSERT INTO product_ingredients (reg_no, ingredient, ingredient_for_total_use)
        VALUES (?, ?, ?)
        """,
        ingredients,
    )
    conn.commit()
    return len(products), len(ingredients)


def import_famic_apply(conn: sqlite3.Connection, paths: list[Path]) -> int:
    cur = conn.cursor()
    valid_regs = {r[0] for r in cur.execute("SELECT reg_no FROM products")}
    batch: list[tuple] = []
    total = 0
    skipped = 0

    def flush() -> None:
        nonlocal batch, total
        if not batch:
            return
        cur.executemany(
            """
            INSERT INTO applications (
              reg_no, use_category, pesticide_type, name, holder_short,
              crop_name, place, target_pest, purpose, dilution, spray_volume,
              timing, times_per_crop, method, fumigation_time, fumigation_temp,
              applicable_soil, zone, applicable_pesticide_name, mix_count,
              ingredient1_total_uses, ingredient2_total_uses, ingredient3_total_uses,
              ingredient4_total_uses, ingredient5_total_uses
            ) VALUES (
              ?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?
            )
            """,
            batch,
        )
        total += len(batch)
        batch = []

    for path in paths:
        _, rows = load_xls_rows(path)
        for row in rows:
            reg = row.get("登録番号", "")
            crop = row.get("作物名", "")
            if not reg or not crop:
                skipped += 1
                continue
            if reg not in valid_regs:
                skipped += 1
                continue
            batch.append(
                (
                    reg,
                    row.get("用途", ""),
                    row.get("農薬の種類", ""),
                    row.get("農薬の名称", ""),
                    row.get("登録を有する者の略称", ""),
                    crop,
                    row.get("適用場所", ""),
                    row.get("適用病害虫雑草名", ""),
                    row.get("使用目的", ""),
                    row.get("希釈倍数使用量", ""),
                    row.get("散布液量", ""),
                    row.get("使用時期", ""),
                    row.get("本剤の使用回数", ""),
                    row.get("使用方法", ""),
                    row.get("くん蒸時間", ""),
                    row.get("くん蒸温度", ""),
                    row.get("適用土壌", ""),
                    row.get("適用地帯名", ""),
                    row.get("適用農薬名", ""),
                    row.get("混合数", ""),
                    row.get("有効成分①を含む農薬の総使用回数", ""),
                    row.get("有効成分②を含む農薬の総使用回数", ""),
                    row.get("有効成分③を含む農薬の総使用回数", ""),
                    row.get("有効成分④を含む農薬の総使用回数", ""),
                    row.get("有効成分⑤を含む農薬の総使用回数", ""),
                )
            )
            if len(batch) >= 1000:
                flush()
                conn.commit()
                print(f"  applications so far: {total}", flush=True)
        flush()
        conn.commit()

    print(f"  applications skipped: {skipped}", flush=True)
    return total


def import_croplife(conn: sqlite3.Connection, path: Path) -> tuple[int, int]:
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet_name = next((s for s in wb.sheetnames if "製品名順" in s), None)
    if not sheet_name:
        raise RuntimeError(f"製品名順シートが見つかりません: {wb.sheetnames}")
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))

    header_i = None
    for i, r in enumerate(rows[:10]):
        if r and any(x and "登録番号" in str(x) for x in r):
            header_i = i
            break
    if header_i is None:
        raise RuntimeError("CropLife ヘッダー行（登録番号）が見つかりません")

    hdr = [cell_str(x) if x is not None else "" for x in rows[header_i]]

    # Headers may use halfwidth: RACｺｰﾄﾞ-1
    def find_col(pred) -> int | None:
        for i, h in enumerate(hdr):
            if pred(h):
                return i
        return None

    i_reg = find_col(lambda h: "登録番号" in h)
    i_name = find_col(lambda h: h == "製品名" or "製品名" in h)
    i_use = find_col(lambda h: h == "用途")
    if i_reg is None:
        raise RuntimeError(f"登録番号列がありません: {hdr}")

    ing_cols = []
    rac_cols = []
    for slot in range(1, 6):
        i_ing = find_col(lambda h, s=slot: h == f"有効成分-{s}" or h.endswith(f"有効成分-{s}"))
        i_rac = find_col(
            lambda h, s=slot: f"RAC" in h and h.endswith(f"-{s}")
        )
        ing_cols.append(i_ing)
        rac_cols.append(i_rac)

    valid_regs = {r[0] for r in conn.execute("SELECT reg_no FROM products")}
    batch: list[tuple] = []
    products_with_rac = 0
    skipped_not_in_famic = 0

    for r in rows[header_i + 1 :]:
        if not r or r[i_reg] is None:
            continue
        reg = cell_str(r[i_reg])
        if not reg:
            continue
        if reg not in valid_regs:
            skipped_not_in_famic += 1
            continue
        pname = cell_str(r[i_name]) if i_name is not None else ""
        use = cell_str(r[i_use]) if i_use is not None else ""
        has_any = False
        for slot in range(1, 6):
            i_ing = ing_cols[slot - 1]
            i_rac = rac_cols[slot - 1]
            ing = cell_str(r[i_ing]) if i_ing is not None and i_ing < len(r) else ""
            rac = cell_str(r[i_rac]) if i_rac is not None and i_rac < len(r) else ""
            if not ing and not rac:
                continue
            batch.append((reg, slot, ing or None, rac or None, pname or None, use or None))
            has_any = True
        if has_any:
            products_with_rac += 1

    conn.executemany(
        """
        INSERT INTO croplife_rac (reg_no, slot, ingredient, rac_code, product_name, use_category)
        VALUES (?, ?, ?, ?, ?, ?)
        """,
        batch,
    )
    conn.commit()
    print(f"  croplife rows: {len(batch)}, products: {products_with_rac}, skipped(not in FAMIC): {skipped_not_in_famic}")
    return products_with_rac, len(batch)


def write_meta(conn: sqlite3.Connection, meta: dict[str, str]) -> None:
    conn.executemany(
        "INSERT OR REPLACE INTO import_meta (key, value) VALUES (?, ?)",
        list(meta.items()),
    )
    conn.commit()


def resolve_wrangler_cmd() -> list[str]:
    """Windows では npx 単体だと CreateProcess に失敗するため .cmd を使う。"""
    local_names = ("wrangler.cmd", "wrangler") if os.name == "nt" else ("wrangler",)
    for name in local_names:
        candidate = ROOT / "node_modules" / ".bin" / name
        if candidate.is_file():
            return [str(candidate)]
    npx = shutil.which("npx.cmd") or shutil.which("npx")
    if not npx:
        raise FileNotFoundError(
            "wrangler / npx が見つかりません。先に npm install を実行してください。"
        )
    return [npx, "wrangler"]


def _row_to_sql_values(row: tuple) -> str:
    vals = []
    for v in row:
        if v is None:
            vals.append("NULL")
        elif isinstance(v, (int, float)):
            vals.append(str(v))
        else:
            vals.append(sql_quote(str(v)))
    return "(" + ", ".join(vals) + ")"


def export_sql_batches(
    conn: sqlite3.Connection,
    out_dir: Path,
    rows_per_stmt: int = 80,
    stmts_per_file: int = 12,
) -> list[Path]:
    """Export INSERT statements for D1 (schema applied separately).

    D1/local SQLite rejects oversized single statements (SQLITE_TOOBIG),
    so keep each INSERT small and group several per file to limit wrangler calls.
    """
    if out_dir.exists():
        for p in out_dir.glob("*.sql"):
            p.unlink()
    out_dir.mkdir(parents=True, exist_ok=True)

    files: list[Path] = []
    # Clear tables then insert (migration already created schema + crop_presets)
    clear_path = out_dir / "00_clear.sql"
    clear_path.write_text(
        "DELETE FROM croplife_rac;\n"
        "DELETE FROM applications;\n"
        "DELETE FROM product_ingredients;\n"
        "DELETE FROM products;\n"
        "DELETE FROM import_meta;\n",
        encoding="utf-8",
    )
    files.append(clear_path)

    def dump_table(table: str, columns: list[str], prefix: str) -> None:
        nonlocal files
        col_list = ", ".join(columns)
        cur = conn.execute(f"SELECT {col_list} FROM {table}")
        part = 0
        stmt_buf: list[str] = []

        def flush() -> None:
            nonlocal part, stmt_buf
            if not stmt_buf:
                return
            part += 1
            path = out_dir / f"{prefix}_{part:04d}.sql"
            path.write_text("".join(stmt_buf), encoding="utf-8")
            files.append(path)
            stmt_buf = []

        while True:
            rows = cur.fetchmany(rows_per_stmt)
            if not rows:
                break
            values = ",\n".join(_row_to_sql_values(row) for row in rows)
            stmt_buf.append(f"INSERT INTO {table} ({col_list}) VALUES\n{values};\n")
            if len(stmt_buf) >= stmts_per_file:
                flush()
        flush()

    dump_table(
        "products",
        [
            "reg_no",
            "pesticide_type",
            "name",
            "holder",
            "concentration",
            "mix_count",
            "use_category",
            "formulation",
            "registered_on",
        ],
        "01_products",
    )
    dump_table(
        "product_ingredients",
        ["reg_no", "ingredient", "ingredient_for_total_use"],
        "02_ingredients",
    )
    # applications: omit id (autoincrement)
    dump_table(
        "applications",
        [
            "reg_no",
            "use_category",
            "pesticide_type",
            "name",
            "holder_short",
            "crop_name",
            "place",
            "target_pest",
            "purpose",
            "dilution",
            "spray_volume",
            "timing",
            "times_per_crop",
            "method",
            "fumigation_time",
            "fumigation_temp",
            "applicable_soil",
            "zone",
            "applicable_pesticide_name",
            "mix_count",
            "ingredient1_total_uses",
            "ingredient2_total_uses",
            "ingredient3_total_uses",
            "ingredient4_total_uses",
            "ingredient5_total_uses",
        ],
        "03_applications",
    )
    dump_table(
        "croplife_rac",
        ["reg_no", "slot", "ingredient", "rac_code", "product_name", "use_category"],
        "04_croplife",
    )
    dump_table("import_meta", ["key", "value"], "05_meta")
    return files


def apply_d1(
    db_name: str,
    sql_files: list[Path],
    apply_schema: bool,
    *,
    remote: bool = False,
) -> None:
    wrangler = resolve_wrangler_cmd()
    target_flag = "--remote" if remote else "--local"

    def run(file_path: Path) -> None:
        args = [
            *wrangler,
            "d1",
            "execute",
            db_name,
            target_flag,
            "--file",
            str(file_path),
            "--yes",
        ]
        print(">", " ".join(args), flush=True)
        # Windows の .cmd 実行では shell=True が必要な場合がある
        use_shell = os.name == "nt" and args[0].lower().endswith(".cmd")
        subprocess.run(
            args if not use_shell else subprocess.list2cmdline(args),
            cwd=str(ROOT),
            check=True,
            shell=use_shell,
        )

    total = len(sql_files) + (1 if apply_schema else 0)
    step = 0
    if apply_schema:
        step += 1
        print(f"[{step}/{total}] schema", flush=True)
        run(MIGRATION)
    for path in sql_files:
        step += 1
        print(f"[{step}/{total}] {path.name}", flush=True)
        run(path)


def apply_d1_local(db_name: str, sql_files: list[Path], apply_schema: bool) -> None:
    apply_d1(db_name, sql_files, apply_schema, remote=False)


def print_summary(conn: sqlite3.Connection) -> None:
    def one(sql: str) -> int:
        return int(conn.execute(sql).fetchone()[0])

    products = one("SELECT COUNT(*) FROM products")
    apps = one("SELECT COUNT(*) FROM applications")
    ings = one("SELECT COUNT(*) FROM product_ingredients")
    rac_products = one("SELECT COUNT(DISTINCT reg_no) FROM croplife_rac")
    matched = one(
        "SELECT COUNT(*) FROM products p WHERE EXISTS (SELECT 1 FROM croplife_rac c WHERE c.reg_no = p.reg_no)"
    )
    print("\n=== import summary ===")
    print(f"products:            {products}")
    print(f"product_ingredients: {ings}")
    print(f"applications:        {apps}")
    print(f"croplife products:   {rac_products}")
    print(f"products with RAC:   {matched} ({100 * matched / products:.1f}%)" if products else "")
    print(f"sqlite:              {SQLITE_PATH}")


def main() -> int:
    parser = argparse.ArgumentParser(description="FAMIC + CropLife → SQLite / D1 import")
    parser.add_argument("--input-dir", help="FAMIC/CropLife ファイルのあるディレクトリ")
    parser.add_argument("--sqlite", type=Path, default=SQLITE_PATH, help="出力 SQLite パス")
    parser.add_argument("--export-sql", action="store_true", default=True, help="D1 用 SQL バッチを出力")
    parser.add_argument("--no-export-sql", action="store_true", help="SQL バッチ出力をしない")
    parser.add_argument("--apply-d1-local", action="store_true", help="ローカル D1 に投入する")
    parser.add_argument("--apply-d1-remote", action="store_true", help="リモート（本番）D1 に投入する")
    parser.add_argument("--d1-name", default=DEFAULT_DB_NAME, help="D1 database name")
    parser.add_argument("--skip-schema", action="store_true", help="D1 投入時に migration をスキップ")
    parser.add_argument(
        "--from-existing-sqlite",
        action="store_true",
        help="既存 SQLite から SQL 再出力して D1 投入のみ（FAMIC 再取込しない）",
    )
    args = parser.parse_args()

    if not MIGRATION.is_file():
        print(f"migration がありません: {MIGRATION}", file=sys.stderr)
        return 1

    sql_files: list[Path] = []
    n_prod = n_app = n_rac_prod = 0

    if args.from_existing_sqlite:
        if not Path(args.sqlite).is_file():
            print(f"SQLite がありません: {args.sqlite}", file=sys.stderr)
            return 1
        print(f"\nusing existing sqlite: {args.sqlite}")
        conn = sqlite3.connect(str(args.sqlite))
        print_summary(conn)
        if args.export_sql and not args.no_export_sql:
            print("\nexporting D1 SQL batches...")
            sql_files = export_sql_batches(conn, SQL_BATCH_DIR)
            print(f"  batches: {len(sql_files)} -> {SQL_BATCH_DIR}")
        conn.close()
    else:
        if not args.input_dir:
            print("--input-dir か --from-existing-sqlite が必要です", file=sys.stderr)
            return 1
        input_dir = Path(args.input_dir)
        if not input_dir.is_dir():
            print(f"入力ディレクトリがありません: {input_dir}", file=sys.stderr)
            return 1

        basic = find_file(input_dir, "登録基本部.xls")
        apply1 = find_file(input_dir, "登録適用部一.xls")
        apply2 = find_file(input_dir, "登録適用部二.xls")
        rac_xlsx = find_file(input_dir, "mechanism_rac.xlsx")

        print("input files:")
        print(f"  basic:  {basic}")
        print(f"  apply1: {apply1}")
        print(f"  apply2: {apply2}")
        print(f"  rac:    {rac_xlsx}")

        print("\nbuilding sqlite...")
        conn = init_sqlite(args.sqlite)

        print("import FAMIC basic...")
        n_prod, n_ing = import_famic_basic(conn, basic)
        print(f"  products={n_prod}, ingredients={n_ing}")

        print("import FAMIC applications...")
        n_app = import_famic_apply(conn, [apply1, apply2])
        print(f"  applications={n_app}")

        print("import CropLife RAC...")
        n_rac_prod, n_rac_rows = import_croplife(conn, rac_xlsx)

        write_meta(
            conn,
            {
                "imported_at": datetime.now(timezone.utc).isoformat(),
                "famic_basic": str(basic),
                "famic_apply1": str(apply1),
                "famic_apply2": str(apply2),
                "croplife_xlsx": str(rac_xlsx),
                "product_count": str(n_prod),
                "ingredient_count": str(n_ing),
                "application_count": str(n_app),
                "croplife_product_count": str(n_rac_prod),
                "croplife_row_count": str(n_rac_rows),
                "schema_version": "0001_pesticides",
            },
        )

        print_summary(conn)

        if args.export_sql and not args.no_export_sql:
            print("\nexporting D1 SQL batches...")
            sql_files = export_sql_batches(conn, SQL_BATCH_DIR)
            print(f"  batches: {len(sql_files)} -> {SQL_BATCH_DIR}")

        conn.close()

        summary_path = OUT_DIR / "last_import_summary.json"
        summary_path.write_text(
            json.dumps(
                {
                    "imported_at": datetime.now(timezone.utc).isoformat(),
                    "sqlite": str(args.sqlite),
                    "products": n_prod,
                    "applications": n_app,
                    "croplife_products": n_rac_prod,
                    "sql_batches": len(sql_files),
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )

    if args.apply_d1_local or args.apply_d1_remote:
        if not sql_files:
            if not Path(args.sqlite).is_file():
                print(f"SQLite がありません: {args.sqlite}", file=sys.stderr)
                return 1
            sql_files = export_sql_batches(sqlite3.connect(str(args.sqlite)), SQL_BATCH_DIR)
        if args.apply_d1_local:
            print("\napplying to local D1...")
            apply_d1(args.d1_name, sql_files, apply_schema=not args.skip_schema, remote=False)
            print("local D1 apply done.")
        if args.apply_d1_remote:
            print("\napplying to remote D1...")
            apply_d1(args.d1_name, sql_files, apply_schema=not args.skip_schema, remote=True)
            print("remote D1 apply done.")

    print("\n完了。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
